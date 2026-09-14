"""Génère le PowerPoint final de la soutenance.

Toutes les valeurs chiffrées sont lues dans results/summary.json, lui-même
produit par `python main.py`. Aucun chiffre n'est écrit en dur.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

OUT_DIR = ROOT / "FINAL_SOUTENANCE"
ASSETS = OUT_DIR / "assets" / "slide"   # figures sans titre interne
OUT_PPTX = OUT_DIR / "Travelling_Merchant_Romain_Yannis_Lisa_FINAL.pptx"

S = json.loads((ROOT / "results" / "summary.json").read_text(encoding="utf-8"))
CHRIST = S["christofides"]["distance_km"]
GA = S["genetic"]["distance_km"]
MST = S["christofides"]["mst_weight_km"]
MATCH = S["christofides"]["matching_weight_km"]
ODD = S["christofides"]["odd_vertices_count"]
GAP = CHRIST - GA
GAP_PCT = 100 * GAP / CHRIST
CFG = S["genetic"]["configuration"]


def fr(v: float, dec: int = 2) -> str:
    return f"{v:,.{dec}f}".replace(",", " ").replace(".", ",")


# --- Charte ---------------------------------------------------------------
BG = RGBColor(0x0B, 0x12, 0x20)
PANEL = RGBColor(0x16, 0x22, 0x3A)
PANEL_HI = RGBColor(0x1E, 0x2E, 0x4D)
LINE = RGBColor(0x25, 0x33, 0x4F)
TEXT = RGBColor(0xF2, 0xF6, 0xFC)
MUTED = RGBColor(0x9D, 0xB0, 0xCE)
GOLD = RGBColor(0xE8, 0xB4, 0x5A)
GOLD_SOFT = RGBColor(0xF5, 0xD0, 0x8A)
CYAN = RGBColor(0x4F, 0xD1, 0xE8)
GREEN = RGBColor(0x5B, 0xD6, 0xA8)
CORAL = RGBColor(0xF2, 0x76, 0x5B)

FONT = "Segoe UI"
W, H = Inches(13.333), Inches(7.5)

# Zone utile pour les visuels
IMG_TOP, IMG_BOTTOM = Inches(1.62), Inches(6.92)
IMG_LEFT, IMG_RIGHT = Inches(0.55), Inches(12.78)


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, *, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         italic=False, space_before=0, space_after=0, first=False, spacing=1.0,
         font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def rect(slide, x, y, w, h, fill=PANEL, line=None, lw=1.25,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BG, shape=MSO_SHAPE.RECTANGLE)
    return slide


def slide_header(slide, eyebrow, title, subtitle=None):
    """Bandeau de titre commun à toutes les slides de contenu."""
    rect(slide, Inches(0.55), Inches(0.46), Inches(0.055), Inches(0.72),
         fill=GOLD, shape=MSO_SHAPE.RECTANGLE)
    tf = textbox(slide, Inches(0.78), Inches(0.40), Inches(11.9), Inches(0.9))
    para(tf, eyebrow.upper(), size=11.5, color=GOLD, bold=True, first=True,
         spacing=0.95)
    para(tf, title, size=27, color=TEXT, bold=True, space_before=1, spacing=0.95)
    if subtitle:
        stf = textbox(slide, Inches(0.78), Inches(1.28), Inches(11.9), Inches(0.35))
        para(stf, subtitle, size=13.5, color=MUTED, first=True)


def footer(slide, number, speaker):
    rect(slide, Inches(0.55), Inches(7.04), Inches(12.23), Pt(0.75),
         fill=LINE, shape=MSO_SHAPE.RECTANGLE)
    tf = textbox(slide, Inches(0.55), Inches(7.12), Inches(8.0), Inches(0.28))
    para(tf, "Le marchand ambulant  ·  TSP  ·  Romain, Yannis & Lisa",
         size=9.5, color=MUTED, first=True)
    tf2 = textbox(slide, Inches(8.8), Inches(7.12), Inches(3.98), Inches(0.28))
    para(tf2, f"{speaker}   ·   {number:02d}", size=9.5, color=MUTED,
         align=PP_ALIGN.RIGHT, first=True)


def place_image(slide, name, box_left, box_top, box_right, box_bottom,
                align="center"):
    """Insère une image en la contenant dans la boîte, sans déformation."""
    path = ASSETS / name
    with Image.open(path) as im:
        iw, ih = im.size
    bw, bh = box_right - box_left, box_bottom - box_top
    scale = min(bw / iw, bh / ih)
    w, h = int(iw * scale), int(ih * scale)
    if align == "center":
        x = box_left + (bw - w) // 2
    elif align == "left":
        x = box_left
    else:
        x = box_right - w
    y = box_top + (bh - h) // 2
    slide.shapes.add_picture(str(path), x, y, Emu(w), Emu(h))
    return x, y, w, h


def stat_card(slide, x, y, w, h, value, label, color=CYAN, vsize=24, lsize=11):
    rect(slide, x, y, w, h, fill=PANEL, line=LINE, lw=1.0)
    rect(slide, x, y, Inches(0.045), h, fill=color, shape=MSO_SHAPE.RECTANGLE)
    tf = textbox(slide, x + Inches(0.24), y + Inches(0.10), w - Inches(0.4),
                 h - Inches(0.2), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, value, size=vsize, color=color, bold=True, first=True, spacing=0.95)
    para(tf, label, size=lsize, color=MUTED, space_before=2, spacing=0.95)


WORDS_PER_MINUTE = 150   # débit courant à l'oral en français, en présentation


def estimate(script: str) -> tuple[int, str]:
    """Durée de lecture du script, arrondie aux 5 secondes supérieures.

    Le temps annoncé est calculé sur le texte réellement écrit plutôt que
    saisi à la main : les deux ne peuvent donc pas diverger.
    """
    seconds = len(script.split()) / WORDS_PER_MINUTE * 60
    seconds = int(-(-seconds // 5) * 5)
    minutes, rest = divmod(seconds, 60)
    if not minutes:
        return seconds, f"{rest} s"
    return seconds, f"{minutes} min" + (f" {rest:02d}" if rest else "")


def notes(slide, speaker, _duration_hint, script, transition):
    """Écrit les notes du présentateur.

    `_duration_hint` est ignoré : la durée est recalculée à partir du script
    (voir `estimate`), pour qu'elle reste exacte quand le texte évolue.
    """
    _, duration = estimate(script)
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"INTERVENANT : {speaker}"
    for line in (f"TEMPS ESTIME : {duration}", "", "CE QU'IL DOIT DIRE :", script,
                 "", f"TRANSITION VERS LA SLIDE SUIVANTE : {transition}"):
        p = tf.add_paragraph()
        p.text = line
